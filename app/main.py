from fastapi import FastAPI, UploadFile, File
from fastapi.responses import FileResponse
from fastapi.staticfiles import StaticFiles

from docx import Document

from openpyxl import load_workbook

import xlrd
from xlutils.copy import copy

import pdfplumber
import os
import re
import json
import traceback

app = FastAPI()

app.mount(
    "/static",
    StaticFiles(directory="app/static"),
    name="static"
)

# =========================
# ROOT
# =========================

@app.get("/")
def root():

    return {
        "status": "vendor-agent-running"
    }


# =========================
# FRONTEND
# =========================

@app.get("/app")
def app_page():

    return FileResponse(
        "app/static/index.html"
    )


# =========================
# EXTRAER TEXTO PDF
# =========================

def extract_text_from_pdf(pdf_path):

    text = ""

    try:

        with pdfplumber.open(pdf_path) as pdf:

            for page in pdf.pages:

                page_text = page.extract_text()

                if page_text:
                    text += page_text + "\n"

    except Exception as e:

        print("PDF ERROR:", e)

    return text


# =========================
# EXTRAER DATOS EMPRESA
# =========================

def extract_company_data(text):

    data = {}

    # =====================
    # NIT
    # =====================

    nit_patterns = [

        r'NIT[:\s]+(\d{6,15})',
        r'(\d{3}[.,]?\d{3}[.,]?\d{3}-?\d)',
    ]

    for pattern in nit_patterns:

        match = re.search(
            pattern,
            text,
            re.IGNORECASE
        )

        if match:

            data["nit"] = match.group(1)
            break

    # =====================
    # EMAIL
    # =====================

    email_match = re.search(
        r'[\w\.-]+@[\w\.-]+\.\w+',
        text
    )

    if email_match:
        data["email"] = email_match.group()

    # =====================
    # TELEFONO
    # =====================

    phone_match = re.search(
        r'(\+57\s?\d{10}|\d{7,10})',
        text
    )

    if phone_match:
        data["telefono"] = phone_match.group()

    # =====================
    # EMPRESA
    # =====================

    empresa_patterns = [

        r'([A-Z\s]+(?:SAS|LTDA|SA|S\.A\.S\.|LIMITADA))',

    ]

    for pattern in empresa_patterns:

        match = re.search(pattern, text)

        if match:

            data["empresa"] = match.group(1).strip()
            break

    # =====================
    # REPRESENTANTE LEGAL
    # =====================

    representante_patterns = [

        r'REPRESENTANTE LEGAL[:\s]+([A-Z\s]+)',
        r'Representante Legal[:\s]+([A-Z\s]+)',

    ]

    for pattern in representante_patterns:

        match = re.search(
            pattern,
            text
        )

        if match:

            data["representante_legal"] = \
                match.group(1).strip()

            break

    # =====================
    # DIRECCION
    # =====================

    direccion_match = re.search(
        r'(CALLE|CARRERA|CRA|CL)\s+[A-Z0-9\s\-#]+',
        text,
        re.IGNORECASE
    )

    if direccion_match:

        data["direccion"] = direccion_match.group()

    return data


# =========================
# SUBIR DOCUMENTOS
# =========================

@app.post("/upload-company-documents")
async def upload_company_documents(
    files: list[UploadFile] = File(...)
):

    try:

        os.makedirs(
            "storage/company_docs",
            exist_ok=True
        )

        all_text = ""

        for file in files:

            file_path = (
                f"storage/company_docs/{file.filename}"
            )

            with open(file_path, "wb") as f:

                f.write(await file.read())

            text = extract_text_from_pdf(file_path)

            all_text += text + "\n"

        company_data = extract_company_data(
            all_text
        )

        with open(
            "storage/company_data.json",
            "w",
            encoding="utf-8"
        ) as json_file:

            json.dump(
                company_data,
                json_file,
                indent=4,
                ensure_ascii=False
            )

        return {

            "status": "success",
            "message": "Documents processed",
            "company_data": company_data,
            "preview_text": all_text[:3000]

        }

    except Exception as e:

        return {

            "status": "error",
            "detail": str(e),
            "trace": traceback.format_exc()

        }


# =========================
# LLENAR WORD
# =========================

@app.post("/fill-word")
async def fill_word(
    file: UploadFile = File(...)
):

    try:

        with open(
            "storage/company_data.json",
            "r",
            encoding="utf-8"
        ) as f:

            company_data = json.load(f)

        os.makedirs(
            "storage/forms",
            exist_ok=True
        )

        input_path = (
            f"storage/forms/{file.filename}"
        )

        with open(input_path, "wb") as f:

            f.write(await file.read())

        doc = Document(input_path)

        replacements = {

            "{{empresa}}":
                company_data.get("empresa", ""),

            "{{nit}}":
                company_data.get("nit", ""),

            "{{representante_legal}}":
                company_data.get(
                    "representante_legal",
                    ""
                ),

            "{{direccion}}":
                company_data.get("direccion", ""),

            "{{telefono}}":
                company_data.get("telefono", ""),

            "{{email}}":
                company_data.get("email", ""),

        }

        for paragraph in doc.paragraphs:

            for key, value in replacements.items():

                if key in paragraph.text:

                    paragraph.text = \
                        paragraph.text.replace(
                            key,
                            value
                        )

        output_path = (
            "storage/FILLED_FORM.docx"
        )

        doc.save(output_path)

        return FileResponse(
            output_path,
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            filename="FILLED_FORM.docx"
        )

    except Exception as e:

        return {

            "status": "error",
            "detail": str(e),
            "trace": traceback.format_exc()

        }


# =========================
# LLENAR EXCEL
# =========================

@app.post("/fill-excel")
async def fill_excel(file: UploadFile = File(...)):

    try:

        os.makedirs("storage", exist_ok=True)

        filename = file.filename.lower()

        input_path = f"storage/{file.filename}"

        with open(input_path, "wb") as f:
            f.write(await file.read())

        if filename.endswith(".xlsx"):

            output_path = "storage/FILLED_EXCEL.xlsx"

            workbook = openpyxl.load_workbook(input_path)

            sheet = workbook.active

            sheet["A1"] = "EMPRESA"
            sheet["B1"] = "ZEHIRUT LTDA"

            workbook.save(output_path)

            return FileResponse(
                path=output_path,
                filename="FILLED_EXCEL.xlsx",
                media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )

        elif filename.endswith(".xls"):

            output_path = "storage/FILLED_EXCEL.xls"

            rb = xlrd.open_workbook(
                input_path,
                formatting_info=True
            )

            wb = copy(rb)

            ws = wb.get_sheet(0)

            ws.write(0, 0, "EMPRESA")
            ws.write(0, 1, "ZEHIRUT LTDA")

            wb.save(output_path)

            return FileResponse(
                path=output_path,
                filename="FILLED_EXCEL.xls",
                media_type="application/vnd.ms-excel"
            )

        else:

            return {
                "error": "Formato no soportado"
            }

    except Exception as e:

        import traceback

        return {
            "error": str(e),
            "trace": traceback.format_exc()
        }
